#!/usr/bin/env python3
"""Tests for the QA checks: (a) on the current outputs, (b) on a fixture pack (two risks, a real
workbook, a real register and pack built by build_register.py / build_pack.py) with planted faults:
a bad figure, a missing row, a wrong quote, a broken trace, an open challenge.

Run: .venv/bin/python scripts/checks/test_checks.py"""
import contextlib
import copy
import io
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1] / ".claude" / "skills" / "audit-program"))
import check_numbers, check_quotes, check_samples, check_trace, run_checks  # noqa: E402
from _util import REPO, safe_run  # noqa: E402
import build_pack as bp  # noqa: E402

br = bp.br
failures = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        failures.append(msg)


# ------------------------------------------------------------------ fixture
ROWS = [(9000, "Acme Ltd"), (9500, "Acme Ltd"), (9800, "Acme Ltd"), (990, "Bolt Ltd"), (980, "Bolt Ltd"),
        (975, "Cog Ltd"), (960, "Cog Ltd"), (2000, "Dyno Ltd"), (3000, "Dyno Ltd"), (500, "Elm Ltd")]
FIRST_DATA_ROW = 3   # sheet row 1 is a title, row 2 the header
ANALYTICS = {"generated_at": "2026-01-01T00:00:00+00:00",
             "limitations": ["Spend dates are payment-run timestamps, not invoice dates."],
             "tests": {"split_purchases@25000": {"counts": {"flagged_items": 3}, "totals_gbp": {"combined": 90000.0},
                                                 "top20": [{"supplier": "ACME"}]},
                       "threshold_clustering@1000": {"counts": {"flagged_items": 4}}}}
HISTORY = {"scoring_method": {"usable": False}, "audits": [{"id": "AUD-01", "audit_title": "Contract Letting",
                                                           "opinion": "Reasonable Assurance"}]}


def risk(rid, title, rules, lik, imp, metric):
    return {"id": rid, "title": title, "theme": rid.lower(),
            "statement": "Risk indicators to investigate " + " ".join(f"[rule:{r}]" for r in rules) + ".",
            "likelihood": {"base_score": lik, "rubric": f"L{lik}",
                           "justification": f"{{{{{metric}.counts.flagged_items}}}} flagged [metric:{metric}.counts.flagged_items].",
                           "uplift": {"applied": False, "audit_ids": [], "reason": "Not a repeat finding [history:AUD-01]."}},
            "impact": {"score": imp, "justification": f"Exposure noted [metric:{metric}.counts.flagged_items]."},
            "evidence": [f"Flagged: {{{{{metric}.counts.flagged_items}}}} [metric:{metric}.counts.flagged_items]"],
            "focus": "Recommend a targeted sample."}


RATINGS = {"risks": [risk("R-01", "Same-supplier payments combining past the quote threshold", ["CPR-13", "CPR-60"], 4, 5,
                          "$.tests.split_purchases@25000"),
                     risk("R-02", "Low-value payments clustering at the lowest quote boundary", ["CPR-13"], 2, 3,
                          "$.tests.threshold_clustering@1000")]}
CHALLENGES = """# Challenges

| # | Type | Risk | Challenge | Citation | Ask | Planner response | Status | Verified |
|---|---|---|---|---|---|---|---|---|
| 1 | sample | R-01 | Sample targets the largest groups only | [metric:$.tests.split_purchases@25000.counts.flagged_items] | Add a random draw | RESPONSE #1: declined | declined | |
| 2 | rating | R-02 | Rating looks low | [rule:CPR-13] | Auditor to consider | RESPONSE #2: for the auditor | for the auditor | |
"""


def build_fixture():
    tmp = tempfile.TemporaryDirectory()
    root = Path(tmp.name)
    o, full, sp = root / "outputs", root / "outputs" / "analytics-full", root / "data" / "spend"
    full.mkdir(parents=True)
    sp.mkdir(parents=True)
    (root / "data" / "contract-rules.pdf").symlink_to(REPO / "data" / "contract-rules.pdf")
    shutil.copy(REPO / "outputs" / "rules.json", o / "rules.json")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data to publish"
    ws.append(["West Berkshire Council expenditure"])
    ws.append(["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"])
    for i, (amt, sup) in enumerate(ROWS):
        ws.append(["Highways", "Supplies", "Works", f"2026-05-{i + 1:02d}", amt, sup])
    wb.save(sp / "Fix_-_published.xlsx")
    spend = ["row_id,service,category,narrative,date,amount,supplier,supplier_norm,tag"]
    for i, (amt, sup) in enumerate(ROWS):
        spend.append(f"Fix_-_published:{FIRST_DATA_ROW + i},Highways,Supplies,Works,2026-05-{i + 1:02d} 10:00:00,{amt},{sup},"
                     f"{sup.upper()},general_procurement")
    (full / "cleaned_spend.csv").write_text("\n".join(spend) + "\n")
    r = lambda k: f"Fix_-_published:{FIRST_DATA_ROW + k}"
    (full / "split_purchases.csv").write_text(
        "tier,supplier,combined_gbp,tag,row_ids\n"
        f"25000,Acme Ltd,28300,general_procurement,{r(0)};{r(1)};{r(2)}\n"
        f"25000,Dyno Ltd,5000,general_procurement,{r(7)};{r(8)}\n")
    (full / "threshold_clustering.csv").write_text(
        "tier,row_id,amount,supplier\n" + "".join(f"1000,{r(k)},{ROWS[k][0]},{ROWS[k][1]}\n" for k in (3, 4, 5, 6)))
    for name, obj in (("analytics", ANALYTICS), ("history", HISTORY), ("risk-ratings", RATINGS)):
        (o / f"{name}.json").write_text(json.dumps(obj))
    reg = lambda extra=None: br.main(
        ["--analytics", str(o / "analytics.json"), "--history", str(o / "history.json"), "--rules", str(o / "rules.json"),
         "--ratings", str(o / "risk-ratings.json"), "--register", str(o / "risk-register.md"),
         "--comments", str(o / "auditor-comments.md"), "--state", str(o / ".register-state.json")])
    with contextlib.redirect_stdout(io.StringIO()):
        assert reg() == 0
        p = o / "auditor-comments.md"
        p.write_text(re.sub(r"(\| R-0\d \|.*?\| )(\|)", lambda m: m.group(0), p.read_text()))
        lines = []
        for line in p.read_text().splitlines():
            c = line.split("|")
            if line.startswith("| R-0"):
                c[4] = " APPROVE "
            lines.append("|".join(c))
        p.write_text("\n".join(lines) + "\n")
        rt = copy.deepcopy(RATINGS)
        rt["revision"] = {"round": 2, "changes": []}
        (o / "risk-ratings.json").write_text(json.dumps(rt))
        assert reg() == 0
    cited = {"CPR-13", "CPR-60", "CPR-61"}
    rules = [x["rule_id"] for x in json.load(open(o / "rules.json"))["rules"]]
    plan = {"register_version": 2,
            "memo": {"objective": "Plan testing of {{$pack.risks_in_scope}} approved risks [rule:CPR-13].",
                     "approach": "Targeted samples from the flagged lists, sized by score."},
            "risks": [
                {"risk_id": "R-01", "controls": [{"id": "C-01", "description": "Quoted for [rule:CPR-60] and not split [rule:CPR-13]."}],
                 "tests": [{"id": "T-01", "control_ids": ["C-01"], "objective": "Check for one underlying requirement.",
                            "step": "Compare invoices [metric:$.tests.split_purchases@25000.counts.flagged_items].",
                            "sample": {"file": "split_purchases.csv", "filter": [{"col": "tier", "op": "==", "value": 25000}],
                                       "sort": [{"col": "combined_gbp", "desc": True}], "n": "by_score"},
                            "pbc": ["Orders and invoices"]}]},
                {"risk_id": "R-02", "controls": [{"id": "C-02", "description": "A quote is sought [rule:CPR-61]."}],
                 "tests": [{"id": "T-02", "control_ids": ["C-02"], "objective": "Confirm a quote was sought.",
                            "step": "Inspect the quote record [rule:CPR-61].",
                            "sample": {"file": "threshold_clustering.csv", "sort": [{"col": "amount", "desc": True}], "n": "by_score"},
                            "pbc": ["Orders and invoices"]}]}],
            "not_tested": [{"rule_ids": [x for x in rules if x not in cited], "reason": "Not testable from spend data."}]}
    (o / "audit-plan.json").write_text(json.dumps(plan))
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        code = bp.main(["--outputs", str(o)])
    assert code == 0, "fixture pack failed to build"
    (o / "challenges.md").write_text(CHALLENGES)
    return tmp, root


def run(mod_run, name, root, final=False):
    return safe_run(name, mod_run, root, final)


def sub(path, old, new, count=1):
    t = path.read_text()
    assert old in t, f"{old!r} not in {path.name}"
    path.write_text(t.replace(old, new, count))


def has(res, needle):
    return any(needle in f for f in res["failures"])


# ------------------------------------------------------ (a) current outputs
res = run(check_quotes.run, "quotes", REPO)
check(not res["failures"] and res["checked"] > 70, f"current outputs: check_quotes passes ({res['checked']} checked)")
for name, mod in (("numbers", check_numbers), ("trace", check_trace), ("samples", check_samples)):
    if not (REPO / "outputs" / "pack-figures.json").exists():
        res = run(mod.run, name, REPO)
        check(has(res, "NOT RUN"), f"current outputs: check_{name} reports NOT RUN while no pack is built (never a silent pass)")

# ------------------------------------------------------------ (b) fixture
tmp, root = build_fixture()
o = root / "outputs"
all_res = {n: run(m.run, n, root, True) for n, m in (("quotes", check_quotes), ("numbers", check_numbers),
                                                      ("trace", check_trace), ("samples", check_samples))}
for n, r in all_res.items():
    check(not r["failures"] and r["checked"] > 0, f"fixture: check_{n} passes ({r['checked']} checked) {r['failures'][:2]}")
check(run_checks.main(["--root", str(root), "--final"]) == 0, "fixture: run_checks --final exits 0 (READY FOR SIGN-OFF)")
(o / "challenges.md").rename(o / "challenges.bak")
r = run(check_trace.run, "trace", root, False)
check(not r["failures"] and r["pending"], "interim: absent challenges.md is PENDING, not a failure")
r = run(check_trace.run, "trace", root, True)
check(has(r, "challenges.md does not exist"), "final: absent challenges.md fails")
(o / "challenges.bak").rename(o / "challenges.md")
tmp.cleanup()


def planted(label, mutate, mod, name, needle, final=True):
    t, rt = build_fixture()
    mutate(rt / "outputs", rt)
    res = run(mod.run, name, rt, final)
    check(has(res, needle), f"{label}: {name} fails with {needle!r}" + ("" if has(res, needle) else f" got {res['failures'][:2]}"))
    t.cleanup()


# ---- numbers
planted("planted bad figure (£ in program)", lambda o, r: sub(o / "audit-program.md", "£33,300", "£33,999"),
        check_numbers, "numbers", "£33,999 is not a value")
planted("planted bad figure (3 digits in memo)", lambda o, r: sub(o / "planning-memo.md", "## Approach", "## Approach\n\nAbout 777 payments were reviewed.\n"),
        check_numbers, "numbers", "777 is not a value")
planted("planted bad figure (matrix count)", lambda o, r: sub(o / "risk-control-matrix.md", "5 transactions", "512 transactions"),
        check_numbers, "numbers", "512 is not a value")
t, rt = build_fixture()
sub(rt / "outputs" / "audit-program.md", "## Approach" if False else "# Audit program", "# Audit program\n\nRow Fix_-_published:5, 2026, cl. 6.16, CPR-13, T-01, page 12, v2.\n")
res = run(check_numbers.run, "numbers", rt, True)
check(not res["failures"], "numbers: ids, year, clause, page and version are skipped")
t.cleanup()

# ---- samples
def drop_spend_row(o, r):
    f = o / "analytics-full" / "cleaned_spend.csv"
    f.write_text("\n".join(l for l in f.read_text().splitlines() if not l.startswith("Fix_-_published:4,")) + "\n")


planted("planted missing row (removed from cleaned_spend)", drop_spend_row, check_samples, "samples", "Fix_-_published:4 is not in cleaned_spend.csv")
planted("planted missing row (id in program not in spend)",
        lambda o, r: sub(o / "audit-program.md", "# Audit program", "# Audit program\n\nSee Fix_-_published:99.\n") if False else
        sub(o / "audit-program.md", "# Audit program", "# Audit program\n\nSee Over500_2026_P09_-_published:99.\n"),
        check_samples, "samples", "Over500_2026_P09_-_published:99 is not in cleaned_spend.csv")
planted("planted sheet-row mismatch (workbook supplier changed)",
        lambda o, r: (lambda wb: (wb.active.cell(row=FIRST_DATA_ROW + 3, column=6, value="Bolt Holdings"), wb.save(r / "data/spend/Fix_-_published.xlsx")))(openpyxl.load_workbook(r / "data/spend/Fix_-_published.xlsx")),
        check_samples, "samples", "workbook supplier 'Bolt Holdings'")
planted("planted amount mismatch (workbook amount changed)",
        lambda o, r: (lambda wb: (wb.active.cell(row=FIRST_DATA_ROW, column=5, value=9001), wb.save(r / "data/spend/Fix_-_published.xlsx")))(openpyxl.load_workbook(r / "data/spend/Fix_-_published.xlsx")),
        check_samples, "samples", "workbook amount 9001")
planted("planted sample row amount changed", lambda o, r: sub(o / "samples" / "T-01.csv", ",9000.0,", ",9100.0,") if ",9000.0," in (o / "samples/T-01.csv").read_text() else sub(o / "samples" / "T-01.csv", ",9000,", ",9100,"),
        check_samples, "samples", "amount 9100")
def past_end(o, r):
    with open(o / "analytics-full" / "cleaned_spend.csv", "a") as fh:
        fh.write("Fix_-_published:500,Highways,Supplies,Works,2026-05-30 10:00:00,100,Zed Ltd,ZED,general_procurement\n")
    sub(o / "audit-program.md", "# Audit program", "# Audit program\n\nSee Fix_-_published:500.\n")


planted("row id in spend but past the end of the workbook", past_end, check_samples, "samples",
        "Fix_-_published:500 is not a data row in the council's workbook")
planted("sample row that points at the workbook header row",
        lambda o, r: [sub(o / f, "Fix_-_published:3,", "Fix_-_published:2,") for f in ("analytics-full/cleaned_spend.csv", "samples/T-01.csv")],
        check_samples, "samples", "Fix_-_published:2 is not a data row in the council's workbook")

# ---- trace
planted("unaccounted rule", lambda o, r: (lambda p: (p["not_tested"][0]["rule_ids"].remove("CPR-02"), (o / "audit-plan.json").write_text(json.dumps(p))))(json.loads((o / "audit-plan.json").read_text())),
        check_trace, "trace", "rule CPR-02: not cited")
planted("not-tested rule without reason", lambda o, r: (lambda p: (p["not_tested"][0].update(reason=""), (o / "audit-plan.json").write_text(json.dumps(p))))(json.loads((o / "audit-plan.json").read_text())),
        check_trace, "trace", "listed as not tested with no reason")
planted("approved risk missing from the program", lambda o, r: sub(o / "audit-program.md", "## R-02:", "## R2:"),
        check_trace, "trace", "R-02: not in audit-program.md")
planted("approved risk missing from the matrix", lambda o, r: sub(o / "risk-control-matrix.md", "| R-02 ", "| R2 "),
        check_trace, "trace", "R-02: not in risk-control-matrix.md")
planted("control with no test", lambda o, r: (lambda p: (p["risks"][1]["tests"][0].update(control_ids=["C-01"]), (o / "audit-plan.json").write_text(json.dumps(p))))(json.loads((o / "audit-plan.json").read_text())),
        check_trace, "trace", "C-02 (R-02): control has no test")
planted("matrix disagrees with the plan", lambda o, r: sub(o / "risk-control-matrix.md", "| CPR-02 | 4.1 | not tested", "| CPR-02 | 4.1 | tested") if "| CPR-02 | 4.1 | not tested" in (o / "risk-control-matrix.md").read_text() else sub(o / "risk-control-matrix.md", "| not tested |", "| tested |"),
        check_trace, "trace", "risk-control-matrix.md shows")
planted("High risk not in scope", lambda o, r: sub(o / "auditor-comments.md", "| R-01 | Same-supplier payments combining past the quote threshold | 20 Very high | APPROVE |  |", "| R-01 | Same-supplier payments combining past the quote threshold | 20 Very high | REJECT |  |"),
        check_trace, "trace", "R-01: Very high risk is not in scope")
planted("exclusion without a reason", lambda o, r: sub(o / "auditor-comments.md", "| APPROVE |  |", "| REJECT |  |", 2),
        check_trace, "trace", "rejected with no reason recorded")
planted("open challenge", lambda o, r: sub(o / "challenges.md", "| declined |", "| open |"),
        check_trace, "trace", "is not closed")
planted("challenge citing no evidence", lambda o, r: sub(o / "challenges.md", "[rule:CPR-13] | Auditor", "none | Auditor"),
        check_trace, "trace", "cites no evidence")
planted("challenge citing an unknown rule", lambda o, r: sub(o / "challenges.md", "[rule:CPR-13] | Auditor", "[rule:CPR-999] | Auditor"),
        check_trace, "trace", "[rule:CPR-999] does not resolve")
t, rt = build_fixture()
sub(rt / "outputs" / "challenges.md", "| declined |", "| open |")
res = run(check_trace.run, "trace", rt, False)
check(not res["failures"] and any("not closed" in p for p in res["pending"]), "interim: an open challenge row is PENDING, not a failure")
t.cleanup()

# ---- quotes
planted("register cites the wrong clause", lambda o, r: sub(o / "risk-register.md", "[rule:CPR-13 (cl. 6.16)]", "[rule:CPR-13 (cl. 9.9)]"),
        check_quotes, "quotes", "printed as cl. 9.9")
planted("register cites a rule that does not exist", lambda o, r: sub(o / "risk-register.md", "[rule:CPR-13 (cl. 6.16)]", "[rule:CPR-998]"),
        check_quotes, "quotes", "[rule:CPR-998] is not in rules.json")
planted("register excerpt not in the rule's quote", lambda o, r: sub(o / "risk-register.md", "## Summary", '> "payments must always be split" [rule:CPR-13 (cl. 6.16)]\n\n## Summary'),
        check_quotes, "quotes", "is not in the verbatim_quote")
def tamper(o, r):
    d = json.loads((o / "rules.json").read_text())
    d["rules"][12]["verbatim_quote"] += " and some words the council never wrote"
    (o / "rules.json").write_text(json.dumps(d))
planted("tampered quote in rules.json", tamper, check_quotes, "quotes", "verify_quotes.py exit 1")
t, rt = build_fixture()
q = json.loads((rt / "outputs/rules.json").read_text())["rules"][12]["verbatim_quote"]
sub(rt / "outputs" / "risk-register.md", "## Summary", f'> "{q[:40]}" [rule:CPR-13 (cl. 6.16)]\n\n## Summary')
res = run(check_quotes.run, "quotes", rt, True)
check(not res["failures"], f"quotes: a true excerpt of the rule's quote passes {res['failures'][:1]}")
t.cleanup()

print()
print("ALL PASS" if not failures else f"{len(failures)} FAILED")
sys.exit(1 if failures else 0)
