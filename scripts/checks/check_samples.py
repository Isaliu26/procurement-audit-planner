#!/usr/bin/env python3
"""Sample check: every transaction in outputs/samples/*.csv, and every spend row id named in
planning-memo.md, risk-control-matrix.md or audit-program.md, exists in
outputs/analytics-full/cleaned_spend.csv AND at its sheet row in the council's workbook
(data/spend/<stem>.xlsx, row id '<stem>:<sheet row>'), with the same supplier and amount.
Usage: check_samples.py [--root DIR]"""
import csv
import re
import sys

import openpyxl

from _util import DOCS, cli, need, result

ROW_ID = re.compile(r"[\w.-]+_published:\d+")  # plus any workbook stem found in data/spend
SHEET = "Data to publish"


class Workbooks:
    def __init__(self, root):
        self.dir, self.cache = root / "data" / "spend", {}

    def rows(self, stem):
        if stem not in self.cache:
            f = need(self.dir / f"{stem}.xlsx", f"workbook {stem}.xlsx")
            rows = list(openpyxl.load_workbook(f, read_only=True, data_only=True)[SHEET].iter_rows(values_only=True))
            hdr = next((i for i, x in enumerate(rows) if x and "Supplier name" in x and "Net amount" in x), None)
            if hdr is None:
                raise ValueError(f"{stem}: header row not found")
            self.cache[stem] = (rows, rows[hdr].index("Supplier name"), rows[hdr].index("Net amount"), hdr)
        return self.cache[stem]

    def cell(self, row_id):
        stem, n = row_id.rsplit(":", 1)
        rows, sc, ac, hdr = self.rows(stem)
        n = int(n)
        if n - 1 <= hdr or n - 1 >= len(rows):
            return None
        row = rows[n - 1]
        return (str(row[sc]).strip() if row[sc] is not None else ""), row[ac]


def run(root, final=False):
    r = result("samples")
    out = root / "outputs"
    spend = {}
    with open(need(out / "analytics-full" / "cleaned_spend.csv"), newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            spend[row["row_id"]] = row
    wbs = Workbooks(root)

    def verify(row_id, where, supplier=None, amount=None):
        r["checked"] += 1
        sp = spend.get(row_id)
        if sp is None:
            r["failures"].append(f"{where}: {row_id} is not in cleaned_spend.csv")
            return
        if supplier is not None and supplier.strip() != sp["supplier"].strip():
            r["failures"].append(f"{where}: {row_id} supplier {supplier!r} differs from cleaned_spend.csv {sp['supplier']!r}")
        if amount is not None and abs(float(amount) - float(sp["amount"])) > 0.005:
            r["failures"].append(f"{where}: {row_id} amount {amount} differs from cleaned_spend.csv {sp['amount']}")
        try:
            wb = wbs.cell(row_id)
        except (ValueError, KeyError) as e:
            r["failures"].append(f"{where}: {row_id} cannot be read from the workbook: {e}")
            return
        if wb is None:
            r["failures"].append(f"{where}: {row_id} is not a data row in the council's workbook")
            return
        if wb[0] != sp["supplier"].strip():
            r["failures"].append(f"{where}: {row_id} workbook supplier {wb[0]!r} differs from cleaned_spend.csv "
                                 f"{sp['supplier']!r}")
        try:
            if abs(float(wb[1]) - float(sp["amount"])) > 0.005:
                r["failures"].append(f"{where}: {row_id} workbook amount {wb[1]} differs from cleaned_spend.csv {sp['amount']}")
        except (TypeError, ValueError):
            r["failures"].append(f"{where}: {row_id} workbook amount {wb[1]!r} is not numeric")

    files = sorted((out / "samples").glob("*.csv")) if (out / "samples").exists() else []
    if not files:
        r["failures"].append("outputs/samples/ holds no sample file: the pack has not been built")
    for f in files:
        with open(f, newline="", encoding="utf-8") as fh:
            for i, row in enumerate(csv.DictReader(fh), 2):
                verify(row["row_id"], f"{f.name}:{i}", row.get("supplier"), row.get("amount"))
    for doc in DOCS:
        text = need(out / doc).read_text(encoding="utf-8")
        stems = [re.escape(p.stem) for p in wbs.dir.glob("*.xlsx")]
        pat = re.compile("|".join([ROW_ID.pattern] + [rf"(?:{s}):\d+" for s in stems]))
        ids = sorted(set(pat.findall(text)))
        for i in ids:
            verify(i, doc)
        r["notes"].append(f"{doc}: {len(ids)} row id(s) named")
    r["notes"].append(f"{len(files)} sample file(s); workbooks opened: {', '.join(sorted(wbs.cache)) or 'none'}")
    return r


if __name__ == "__main__":
    sys.exit(cli("samples", run))
