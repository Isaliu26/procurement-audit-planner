import sys, os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import glob, datetime as dt, re
import openpyxl, pandas as pd
from config import *
from common import normalise_supplier, sha256
from tagging import tag_row

class InputError(Exception): pass

def read_workbooks():
    files = sorted(glob.glob(f"{ROOT}/data/spend/*.xlsx"))
    if not files: raise InputError("no data/spend/*.xlsx")
    recs, inputs, blank_total = [], [], 0
    for f in files:
        stem = os.path.splitext(os.path.basename(f))[0]
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        if SHEET not in wb.sheetnames: raise InputError(f"{stem}: sheet '{SHEET}' missing")
        rows = list(wb[SHEET].iter_rows(values_only=True))
        hdr = None
        for i, r in enumerate(rows):
            vals = [str(v).strip() if v is not None else "" for v in r]
            if all(c in vals for c in EXPECTED_COLS): hdr = i; idx = [vals.index(c) for c in EXPECTED_COLS]; break
        if hdr is None: raise InputError(f"{stem}: header row with all six columns not found")
        blank = n = 0
        for j in range(hdr + 1, len(rows)):
            r = rows[j]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in r): blank += 1; continue
            svc, cat, nar, d, amt, sup = [r[k] for k in idx]
            if isinstance(d, dt.datetime): pass
            elif isinstance(d, str):
                try: d = dt.datetime.strptime(d.strip(), "%m/%d/%y")
                except ValueError: raise InputError(f"{stem}:{j+1} date text not M/D/YY: {d!r}")
            else: raise InputError(f"{stem}:{j+1} date neither datetime nor text: {d!r}")
            if isinstance(amt, bool) or not isinstance(amt, (int, float)): raise InputError(f"{stem}:{j+1} amount not numeric: {amt!r}")
            recs.append(dict(row_id=f"{stem}:{j+1}", workbook=stem, service=svc, category=cat, narrative=nar, date=d, amount=float(amt), supplier=sup))
            n += 1
        blank_total += blank
        inputs.append({"file": f"data/spend/{os.path.basename(f)}", "sha256": sha256(f), "header_row": hdr + 1, "rows_read": n, "blank_rows_skipped": blank})
    return pd.DataFrame(recs), inputs, blank_total

def classify(df):
    df = df.copy()
    df["supplier_norm"] = df["supplier"].map(normalise_supplier)
    df["tag"] = [tag_row(a, b, c, d) for a, b, c, d in zip(df.narrative, df.category, df.service, df.supplier)]
    reason = []
    for nar, sup, sn in zip(df.narrative, df.supplier, df.supplier_norm):
        s = str(sup or "").strip()
        if s.casefold() == "redacted": reason.append("redacted_supplier")
        elif PENSION_PAT.search(str(nar or "")): reason.append("pension_statutory")
        elif sn == COUNCIL_NAME_NORM: reason.append("internal")
        elif PUBLIC_BODY_PAT.search(s): reason.append("other_public_body")
        else: reason.append("")
    df["excl_reason"] = reason
    return df

def window(df):
    a, b = df.date.min(), df.date.max()
    months = (b.year - a.year) * 12 + b.month - a.month + 1
    return a, b, months
